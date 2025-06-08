import os
from typing import List, Optional, Callable, Dict, Any
import logging
import yaml

try:
    import requests
    HAS_REQUESTS = True
except ImportError:
    HAS_REQUESTS = False
    requests = None

# Optional imports with fallbacks
try:
    from .utils import is_valid_filetype, validate_file
    HAS_UTILS = True
except ImportError:
    HAS_UTILS = False
    def is_valid_filetype(filepath):
        return os.path.exists(filepath)
    def validate_file(filepath):
        return os.path.exists(filepath)

try:
    from PyPDF2 import PdfReader
    HAS_PYPDF2 = True
except ImportError:
    HAS_PYPDF2 = False
    PdfReader = None

try:
    from PIL import Image
    import pytesseract
    HAS_OCR = True
except ImportError:
    HAS_OCR = False
    Image = None
    pytesseract = None

try:
    from .ocr import ocr_file, ocr_pdf_page
    HAS_OCR_MODULE = True
except ImportError:
    HAS_OCR_MODULE = False
    def ocr_file(filepath):
        return f"OCR not available for {filepath}"
    def ocr_pdf_page(pdf_path, page_num):
        return f"OCR not available for {pdf_path} page {page_num}"

try:
    from transformers import AutoModel, AutoTokenizer
    from huggingface_hub import hf_hub_download
    HAS_TRANSFORMERS = True
except ImportError:
    HAS_TRANSFORMERS = False
    AutoModel = None
    AutoTokenizer = None
    hf_hub_download = None

try:
    # Import minimal RAG system for testing
    from .rag_system_minimal import RAGSystem
    HAS_RAG_SYSTEM = True
except ImportError:
    HAS_RAG_SYSTEM = False
    RAGSystem = None

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False
    docx = None

try:
    from sentence_transformers import SentenceTransformer
    HAS_SENTENCE_TRANSFORMERS = True
except ImportError:
    HAS_SENTENCE_TRANSFORMERS = False

# --- RAG Pipeline Core Functions ---
# These functions handle document chunking, prompt loading, and RAG orchestration.

# 2. Set default chunk size/overlap and logging config
DEFAULT_CHUNK_SIZE = 1000
DEFAULT_CHUNK_OVERLAP = 100

LOG_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../.logs'))
if not os.path.exists(LOG_DIR):
    try:
        os.makedirs(LOG_DIR, exist_ok=True)
    except Exception:
        LOG_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '.logs'))
        os.makedirs(LOG_DIR, exist_ok=True)
LOG_FILE = os.path.join(LOG_DIR, 'rag_pipeline.log')

logging.basicConfig(filename=LOG_FILE, level=logging.INFO, format='%(asctime)s %(levelname)s %(message)s')

# 3. Default policy prompt
DEFAULT_POLICY_PROMPT = """You are a helpful assistant. Provide a concise and accurate response based on the context.\n\nContext:\n{context}\n\nQuestion:\n{query}\n\nAnswer:"""

# 4. Load policy prompt from YAML (already present, just ensure import and fallback)
def load_policy_prompt(yaml_path=None):
    """
    Load the policy prompt from policy_prompt.yaml for RAG/LLM continuity.
    If not found or invalid, falls back to DEFAULT_POLICY_PROMPT.
    """
    if yaml_path is None:
        base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../config'))
        yaml_path = os.path.join(base_dir, 'policy_prompt.yaml')
    try:
        with open(yaml_path, 'r', encoding='utf-8') as f:
            data = yaml.safe_load(f)
            if isinstance(data, dict) and 'prompt' in data:
                return data['prompt']
            elif isinstance(data, str):
                return data
            else:
                raise ValueError('policy_prompt.yaml must contain a string or a dict with a "prompt" key.')
    except Exception as e:
        logging.warning(f"Could not load policy prompt from {yaml_path}: {e}. Using default.")
        return DEFAULT_POLICY_PROMPT

# 5. Chunking with overlap

def chunk_text(text: str, chunk_size: int = DEFAULT_CHUNK_SIZE, overlap: int = DEFAULT_CHUNK_OVERLAP) -> List[str]:
    """Chunk text with overlap."""
    chunks = []
    i = 0
    while i < len(text):
        chunk = text[i:i+chunk_size].strip()
        if chunk:
            chunks.append(chunk)
        i += chunk_size - overlap
    return chunks

# 6. Document chunking with OCR for images and image-based PDFs

def chunk_document(document_path: str, chunk_size: int = None, overlap: int = None, config=None) -> List[str]:
    """
    Split a document (PDF, DOCX, TXT, or image) into text chunks for LLM evaluation.
    Returns a list of text chunks.
    Automatically applies OCR for image files and image-based PDFs.
    """
    # Use config-driven settings
    if config is None:
        try:
            import sys
            # Add src directory to path for config import
            src_path = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
            if src_path not in sys.path:
                sys.path.insert(0, src_path)
            from config import load_config
            config = load_config()
        except Exception as e:
            # Fallback to default values if config loading fails
            config = {
                'chunk_size': DEFAULT_CHUNK_SIZE,
                'chunk_overlap': DEFAULT_CHUNK_OVERLAP
            }
    if chunk_size is None:
        chunk_size = config.get('chunk_size', 500)
    if overlap is None:
        overlap = config.get('chunk_overlap', 50)

    ext = os.path.splitext(document_path)[1].lower()
    text = ""
    image_exts = {'.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.tif', '.gif'}
    
    if ext == '.pdf' and HAS_PYPDF2:
        reader = PdfReader(document_path)
        page_texts = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text() or ''
            if not page_text.strip() and HAS_OCR_MODULE:
                # OCR fallback for image-based PDF page
                try:
                    page_text = ocr_pdf_page(document_path, i)
                except Exception as e:
                    logging.warning(f"OCR failed for PDF page {i+1}: {e}")
            page_texts.append(page_text)
        text = "\n".join(page_texts)
    elif ext == '.docx' and HAS_DOCX:
        doc = docx.Document(document_path)
        text = "\n".join(para.text for para in doc.paragraphs)
    elif ext == '.txt':
        with open(document_path, encoding='utf-8', errors='ignore') as f:
            text = f.read()
    elif ext in image_exts and HAS_OCR_MODULE:
        text = ocr_file(document_path)
    elif ext == '.pdf' and not HAS_PYPDF2:
        text = f"PDF processing not available - PyPDF2 missing for {document_path}"
    elif ext == '.docx' and not HAS_DOCX:
        text = f"DOCX processing not available - python-docx missing for {document_path}"
    elif ext in image_exts and not HAS_OCR_MODULE:
        text = f"OCR processing not available - OCR dependencies missing for {document_path}"
    else:
        raise ValueError(f"Unsupported file type for chunking: {ext}")
    return chunk_text(text, chunk_size, overlap)

# 7. Batch processing for chunking

def batch_chunk_documents(file_paths: List[str], chunk_size: int = DEFAULT_CHUNK_SIZE, overlap: int = DEFAULT_CHUNK_OVERLAP) -> List[List[str]]:
    """Batch process multiple documents for chunking."""
    all_chunks = []
    for path in file_paths:
        try:
            if HAS_UTILS:
                validate_file(path)
            chunks = chunk_document(path, chunk_size, overlap)
            all_chunks.append(chunks)
        except Exception as e:
            logging.warning(f"Failed to process {path}: {e}")
    return all_chunks

# --- Extensibility: Pluggable Retriever/LLM ---
def process_with_retriever_and_llm(
    chunks: List[str],
    query: str,
    retriever: Optional[Callable[[str, List[str]], List[str]]] = None,
    llm: Optional[Callable[[str, str], str]] = None,
    policy_prompt: Optional[str] = None,
    top_k: int = None,
    all_chunks: bool = False,
    config=None,
    use_enhanced_rag: bool = True
) -> str:
    """
    Retrieve relevant chunks and pass to LLM for categorization/summarization.
    If all_chunks is True, process all chunks regardless of query relevance.
    If use_enhanced_rag is True, use the new RAG system for better results.
    """
    # Use config-driven settings
    if config is None:
        try:
            from .config import load_config
            config = load_config()
        except Exception:
            config = {'rag_top_k': 5}
    if top_k is None:
        top_k = config.get('rag_top_k', 5)
    if policy_prompt is None:
        policy_prompt = load_policy_prompt()
    
    if all_chunks:
        relevant_chunks = chunks
        metadata = {'method': 'all_chunks'}
    else:
        # Try enhanced RAG first if enabled
        if use_enhanced_rag:
            try:
                query_result = enhanced_query_documents(query, chunks, config, top_k)
                if query_result['success']:
                    relevant_chunks = query_result['relevant_chunks']
                    metadata = query_result['metadata']
                else:
                    # Fallback to legacy retrieval
                    raise Exception("Enhanced RAG failed, falling back to legacy")
            except Exception as e:
                logging.warning(f"Enhanced RAG failed, using legacy retrieval: {e}")
                use_enhanced_rag = False
        
        if not use_enhanced_rag:
            # Legacy retrieval (simple keyword search if no retriever)
            if retriever:
                relevant_chunks = retriever(query, chunks)
                metadata = {'method': 'custom_retriever'}
            else:
                words = set(query.lower().split())
                scored = [(sum(w in c.lower() for w in words), c) for c in chunks]
                scored.sort(reverse=True)
                relevant_chunks = [c for score, c in scored[:top_k] if score > 0]
                if not relevant_chunks:
                    relevant_chunks = chunks[:top_k]
                metadata = {'method': 'keyword_search'}
    
    context = "\n---\n".join(relevant_chunks)
    
    # Create a proper prompt with context and query
    if "{context}" in policy_prompt and "{query}" in policy_prompt:
        prompt = policy_prompt.format(context=context, query=query)
    else:
        # Fallback if policy prompt doesn't have placeholders
        prompt = f"{policy_prompt}\n\nContext:\n{context}\n\nQuery: {query}"
    
    # LLM call
    if llm:
        result = llm(prompt, context)
    else:
        # Try to use the built-in LLM analysis if available
        try:
            from .llm import analyze_chunks
            llm_outputs = analyze_chunks([prompt], provider=config.get('llm_provider', 'koboldcpp'))
            if llm_outputs and llm_outputs[0]:
                result = llm_outputs[0].get('output', str(llm_outputs[0]))
            else:
                result = f"You are a helpful assistant. Provide a concise and accurate response based on the context.\n\nContext:\n{context[:1000]}...\n\nQuery: {query}"
        except Exception as e:
            result = f"You are a helpful assistant. Provide a concise and accurate response based on the context.\n\nContext:\n{context[:1000]}...\n\nQuery: {query}"
    
    # Add metadata to result if possible
    if isinstance(result, dict):
        result['retrieval_metadata'] = metadata
    
    return result

# --- Example Test Function ---
def test_rag_pipeline():
    """Test the RAG pipeline with a sample document and query."""
    import tempfile
    sample_text = "This is a test document.\nIt has several lines.\nThe quick brown fox jumps over the lazy dog."
    with tempfile.NamedTemporaryFile('w+', suffix='.txt', delete=False) as f:
        f.write(sample_text)
        f.flush()
        path = f.name
    try:
        from .utils import validate_file
        validate_file(path)
        chunks = chunk_document(path)
        assert len(chunks) > 0
        result = process_with_retriever_and_llm(chunks, "What does the document say?", policy_prompt="Summarize:")
        print("Test result:", result)
    finally:
        os.remove(path)

def ensure_hf_prefix(model_name):
    """Ensure the model string starts with 'hf.co/' if it looks like a HuggingFace repo."""
    if not isinstance(model_name, str):
        return model_name
    if model_name.startswith('hf.co/'):
        return model_name
    if '/' in model_name:
        return f"hf.co/{model_name}"
    return model_name

def download_with_progress(url, dest_path):
    if not HAS_REQUESTS:
        raise ImportError("requests package required for downloading")
    response = requests.get(url, stream=True)
    total = int(response.headers.get('content-length', 0))
    downloaded = 0
    chunk_size = 8192
    with open(dest_path, 'wb') as f:
        for chunk in response.iter_content(chunk_size=chunk_size):
            if chunk:
                f.write(chunk)
                downloaded += len(chunk)
                percent = int(downloaded * 100 / total) if total else 0
                bar = ('#' * (percent // 2)).ljust(50)
                print(f"\r[DOWNLOAD] |{bar}| {percent}%", end='')
    print(f"\n[INFO] Download complete: {dest_path}")

def ensure_rag_model_downloaded(rag_model, model_dir):
    import os
    if not rag_model:
        return
    # Accept both 'hf.co/author/repo' and 'author/repo' formats
    if rag_model.startswith('hf.co/'):
        model_str = rag_model[len('hf.co/'):]
    else:
        model_str = rag_model
    # Only handle HuggingFace models
    if '/' not in model_str:
        return
    repo_id = model_str
    local_dir = model_dir
    os.makedirs(local_dir, exist_ok=True)
    # ONNX model support: use SentenceTransformer for loading
    if 'onnx' in rag_model.lower():
        if HAS_SENTENCE_TRANSFORMERS:
            try:
                # Use the already imported SentenceTransformer if available
                from sentence_transformers import SentenceTransformer
                print(f"[INFO] Downloading/loading ONNX RAG model: {repo_id}")
                model = SentenceTransformer(rag_model)
                print(f"[INFO] Downloaded and loaded ONNX RAG model: {repo_id}")
                return model
            except Exception as e:
                print(f"[ERROR] Could not download or load ONNX RAG model {repo_id}: {e}")
                return None
        else:
            print(f"[WARNING] sentence_transformers not available for ONNX RAG model {repo_id}")
            return None
    # Default: PyTorch model
    try:
        if HAS_TRANSFORMERS and AutoModel and AutoTokenizer:
            AutoModel.from_pretrained(repo_id, cache_dir=local_dir)
            AutoTokenizer.from_pretrained(repo_id, cache_dir=local_dir)
            print(f"[INFO] Downloaded RAG model: {repo_id}")
        else:
            print(f"[WARNING] Transformers not available - cannot download model {repo_id}")
            return None
    except Exception as e:
        print(f"[ERROR] Could not download RAG model {repo_id}: {e}")
        return None

def get_rag_model(config):
    # Use rag_model if set, otherwise fall back to llm_model
    rag_model = config.get('rag_model', '')
    if not rag_model:
        rag_model = config.get('llm_model', '')
    rag_model = ensure_hf_prefix(rag_model)
    model_dir = config.get('model_dir', 'models')
    ensure_rag_model_downloaded(rag_model, model_dir)
    return rag_model

# --- Enhanced RAG Pipeline with New RAG System ---
# These functions provide modern RAG capabilities using the new RAG system

# Global RAG system instance (lazy initialization)
_rag_system = None

def get_rag_system(config=None):
    """Get or initialize the RAG system instance"""
    global _rag_system
    if _rag_system is None:
        if config is None:
            try:
                from .config import load_config
                config = load_config()
            except Exception:
                # Fallback config
                config = {
                    'rag': {
                        'mode': 'embedded',
                        'embedded': {
                            'model': 'auto',
                            'chunk_size': 1000,
                            'chunk_overlap': 100,
                            'top_k': 5
                        }
                    }
                }
        _rag_system = RAGSystem(config)
    return _rag_system

def enhanced_chunk_document(document_path: str, config=None) -> Dict[str, Any]:
    """
    Enhanced document chunking using the new RAG system.
    Returns both chunks and metadata about the chunking process.
    """
    try:
        rag_system = get_rag_system(config)
        
        # Read the document content
        with open(document_path, 'rb') as f:
            content = f.read()
        
        # Get chunks and metadata
        result = rag_system.chunk_documents([content], [document_path])
        
        if result and len(result) > 0:
            return {
                'chunks': result[0]['chunks'],
                'metadata': result[0]['metadata'],
                'success': True
            }
        else:
            # Fallback to legacy chunking
            legacy_chunks = chunk_document(document_path, config=config)
            return {
                'chunks': legacy_chunks,
                'metadata': {'method': 'legacy_fallback'},
                'success': True
            }
    except Exception as e:
        logging.error(f"Enhanced document chunking failed for {document_path}: {e}")
        # Fallback to legacy method
        try:
            legacy_chunks = chunk_document(document_path, config=config)
            return {
                'chunks': legacy_chunks,
                'metadata': {'method': 'legacy_fallback', 'error': str(e)},
                'success': True
            }
        except Exception as legacy_error:
            return {
                'chunks': [],
                'metadata': {'error': str(legacy_error)},
                'success': False
            }

def enhanced_query_documents(query: str, chunks: List[str], config=None, top_k: int = None) -> Dict[str, Any]:
    """
    Enhanced document querying using the new RAG system.
    Returns relevant chunks with similarity scores and metadata.
    """
    try:
        rag_system = get_rag_system(config)
        
        # Use config defaults if not specified
        if config is None:
            try:
                from .config import load_config
                config = load_config()
            except Exception:
                config = {'rag': {'embedded': {'top_k': 5}}}
        
        if top_k is None:
            top_k = config.get('rag', {}).get('embedded', {}).get('top_k', 5)
        
        # Query the RAG system
        result = rag_system.query(query, chunks, top_k=top_k)
        
        return {
            'relevant_chunks': result.get('chunks', chunks[:top_k]),
            'scores': result.get('scores', []),
            'metadata': result.get('metadata', {}),
            'success': True
        }
    except Exception as e:
        logging.error(f"Enhanced query failed: {e}")
        # Fallback to simple keyword matching
        words = set(query.lower().split())
        scored = [(sum(w in c.lower() for w in words), c) for c in chunks]
        scored.sort(reverse=True)
        relevant_chunks = [c for score, c in scored[:top_k] if score > 0]
        if not relevant_chunks:
            relevant_chunks = chunks[:top_k]
        
        return {
            'relevant_chunks': relevant_chunks,
            'scores': [],
            'metadata': {'method': 'keyword_fallback', 'error': str(e)},
            'success': True
        }

def enhanced_rag_pipeline(document_paths: List[str], query: str, config=None) -> Dict[str, Any]:
    """
    Complete enhanced RAG pipeline using the new RAG system.
    Processes documents, chunks them, and returns query results.
    """
    try:
        rag_system = get_rag_system(config)
        
        # Read all documents
        documents = []
        file_paths = []
        for path in document_paths:
            try:
                validate_file(path)
                with open(path, 'rb') as f:
                    documents.append(f.read())
                file_paths.append(path)
            except Exception as e:
                logging.warning(f"Could not read document {path}: {e}")
                continue
        
        if not documents:
            return {
                'success': False,
                'error': 'No valid documents could be processed',
                'relevant_chunks': [],
                'metadata': {}
            }
        
        # Process documents through RAG system
        chunk_results = rag_system.chunk_documents(documents, file_paths)
        
        # Combine all chunks
        all_chunks = []
        all_metadata = []
        for result in chunk_results:
            all_chunks.extend(result['chunks'])
            all_metadata.append(result['metadata'])
        
        if not all_chunks:
            return {
                'success': False,
                'error': 'No chunks could be extracted from documents',
                'relevant_chunks': [],
                'metadata': {'chunk_metadata': all_metadata}
            }
        
        # Query the chunks
        query_result = rag_system.query(query, all_chunks)
        
        return {
            'success': True,
            'relevant_chunks': query_result.get('chunks', []),
            'scores': query_result.get('scores', []),
            'metadata': {
                'chunk_metadata': all_metadata,
                'query_metadata': query_result.get('metadata', {}),
                'total_chunks': len(all_chunks),
                'documents_processed': len(chunk_results)
            }
        }
        
    except Exception as e:
        logging.error(f"Enhanced RAG pipeline failed: {e}")
        return {
            'success': False,
            'error': str(e),
            'relevant_chunks': [],
            'metadata': {'error': str(e)}
        }

# --- Legacy Compatibility Functions ---
# Maintain backward compatibility while providing enhanced functionality
